import streamlit as st
import os
from PyPDF2 import PdfReader
import docx
from sqlalchemy import create_engine, Column, Integer, String, Text
from sqlalchemy.ext.declarative import declarative_base
from sqlalchemy.orm import sessionmaker
from pptx import Presentation
import openai
from io import BytesIO

# Set OpenAI API Key
openai.api_key = "YOUR_OPENAI_API_KEY"

# Database setup
DATABASE_URL = "sqlite:///due_diligence.db"
engine = create_engine(DATABASE_URL, echo=True)
Base = declarative_base()
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)

# Create models
class Document(Base):
    __tablename__ = 'documents'
    id = Column(Integer, primary_key=True, index=True)
    file_name = Column(String, index=True)
    content = Column(Text)

class Answer(Base):
    __tablename__ = 'answers'
    id = Column(Integer, primary_key=True, index=True)
    document_id = Column(Integer, index=True)
    question = Column(String)
    answer = Column(Text)
    tag = Column(String)

# Create tables
Base.metadata.create_all(bind=engine)

# Helper functions for parsing documents
def parse_pdf(file):
    reader = PdfReader(file)
    text = ""
    for page in reader.pages:
        text += page.extract_text()
    return text

def parse_word(file):
    doc = docx.Document(file)
    text = ""
    for para in doc.paragraphs:
        text += para.text
    return text

# Saving document to database
def save_to_db(file_name, content):
    db = SessionLocal()
    db_document = Document(file_name=file_name, content=content)
    db.add(db_document)
    db.commit()
    db.refresh(db_document)
    return db_document

# Saving answer to database
def save_answer_to_db(document_id, question, answer):
    db = SessionLocal()
    db_answer = Answer(document_id=document_id, question=question, answer=answer)
    db.add(db_answer)
    db.commit()
    db.refresh(db_answer)
    return db_answer

# Generate report (PowerPoint)
def generate_pptx_report(answers):
    prs = Presentation()
    slide_layout = prs.slide_layouts[1]  # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.shapes.placeholders[1]
    
    title.text = "Due Diligence Q&A Report"
    
    # Add Q&A content to slides
    for answer in answers:
        content.text += f"Question: {answer.question}\nAnswer: {answer.answer}\n\n"
    
    pptx_file = "due_diligence_report.pptx"
    prs.save(pptx_file)
    return pptx_file

# Classify answers based on simple keywords
def classify_answer(answer):
    # Example: tagging based on keywords
    if "risk" in answer.lower():
        return "Risk"
    elif "growth" in answer.lower():
        return "Growth"
    else:
        return "General"

# Streamlit UI
st.set_page_config(page_title="Due Diligence Platform", layout="wide")

# Navigation
page = st.sidebar.selectbox("Select Page", ["Home", "About", "Upload Documents", "Generate Report"])

# Home Page
if page == "Home":
    st.markdown("""
        <div style="padding: 4rem 2rem 2rem; background: linear-gradient(to right, #1a1a1a, #0f0f0f); text-align: center;">
            <h1 style="color: white; font-size: 3rem;">Smart Insights.<br>For Your Trusted Decisions</h1>
            <p style="color: #ccc; font-size: 1.2rem;">Streamlined due diligence to empower investors, startups, and enterprises.</p>
        </div>
    """, unsafe_allow_html=True)

# About Page
elif page == "About":
    st.subheader("About Our Due Diligence Platform")
    st.markdown("""
        Our platform streamlines due diligence by automating document parsing, answering questions, 
        generating reports, and enabling collaborative interactions with a Human-in-the-Loop feature.
    """)

# Upload Documents Page
elif page == "Upload Documents":
    uploaded_file = st.file_uploader("Upload your documents", type=["pdf", "docx", "txt"])
    if uploaded_file is not None:
        # Handle document parsing
        if uploaded_file.type == "application/pdf":
            document_text = parse_pdf(uploaded_file)
        elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            document_text = parse_word(uploaded_file)
        else:
            document_text = uploaded_file.read().decode("utf-8")
        
        # Save parsed content to database
        document_saved = save_to_db(uploaded_file.name, document_text)
        st.write(f"Document {document_saved.file_name} saved to database!")

        st.write("Parsed Document Content:")
        st.text_area("Document Content", document_text, height=300)

        # Ask a question
        question = st.text_input("Enter a question to ask")
        if question:
            # Get LLM response (OpenAI example)
            response = openai.Completion.create(
                model="text-davinci-003",
                prompt=f"Answer the following question based on the document: {document_text}\nQuestion: {question}",
                max_tokens=500
            )
            answer = response.choices[0].text.strip()

            # Save answer to database
            answer_saved = save_answer_to_db(document_saved.id, question, answer)
            st.write(f"Answer for '{question}' saved!")
            st.write("Answer:", answer)

# Generate Report Page
elif page == "Generate Report":
    db = SessionLocal()
    answers = db.query(Answer).all()

    if st.button("Generate PowerPoint Report"):
        pptx_file = generate_pptx_report(answers)
        st.write("Report Generated: [Download PPTX Report](%s)" % pptx_file)

    # Display all answers
    for answer in answers:
        st.write(f"Question: {answer.question}")
        st.write(f"Answer: {answer.answer}")
        st.write(f"Tag: {answer.tag}")

    # Tag classification for answers
    if st.button("Classify Answers"):
        for answer in answers:
            answer.tag = classify_answer(answer.answer)
            db.commit()
        st.write("Answers classified successfully.")

# Human in the Loop - Chatbot Interaction
human_input = st.text_input("Improve the answer (Human in the Loop)")
if human_input:
    # Process with LLM or a simple chatbot algorithm
    improved_answer = openai.Completion.create(
        model="text-davinci-003",
        prompt=f"Improve the following answer: {human_input}",
        max_tokens=500
    ).choices[0].text.strip()
    
    st.write("Improved Answer:", improved_answer)
