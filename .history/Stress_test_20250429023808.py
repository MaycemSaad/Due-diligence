import streamlit as st
import pandas as pd
import numpy as np
import plotly.graph_objects as go
from sklearn.preprocessing import MinMaxScaler
from sklearn.linear_model import LogisticRegression
from sklearn.ensemble import IsolationForest
from statsmodels.tsa.arima.model import ARIMA
from prophet import Prophet
from pptx import Presentation
from pptx.util import Inches, Pt
from nltk.sentiment import SentimentIntensityAnalyzer
import requests
from bs4 import BeautifulSoup
import io
import os
import datetime
import nltk
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE




# ========== CONFIGURATION ========== #
st.set_page_config(page_title="Crypto Stress Test AI - Ultimate Pro", layout="wide")

if "mode" not in st.session_state:
    st.session_state.mode = "Light"
mode_choice = st.toggle("🌗 Dark / Light Mode", value=(st.session_state.mode == "Dark"))
st.session_state.mode = "Dark" if mode_choice else "Light"
template = "plotly_dark" if st.session_state.mode == "Dark" else "plotly_white"

FOLDER_PATH = "C:/Users/informasud/Desktop/Data-Preparation/LOT3"

# ========== FUNCTIONS ========== #
try:
    nltk.data.find('sentiment/vader_lexicon')
except LookupError:
    nltk.download('vader_lexicon')
@st.cache_data
def load_files(folder_path):
    return [f for f in os.listdir(folder_path) if f.endswith(".csv")]

@st.cache_data
def load_crypto_data(filepath, lookback_days):
    df = pd.read_csv(filepath)
    df.columns = [col.strip().replace('"', '') for col in df.columns]
    df['Date'] = pd.to_datetime(df['Date'], errors='coerce')
    df = df.dropna(subset=['Date'])
    df['Price'] = df['Price'].astype(str).str.replace(',', '').astype(float)
    df = df.sort_values('Date').set_index('Date')
    df = df[['Price']].rename(columns={"Price": "CRYPTO"})
    df = df.last(f'{lookback_days}D')
    return df

def fit_arima(series, steps):
    model = ARIMA(series, order=(2,1,2))
    fitted = model.fit()
    forecast = fitted.get_forecast(steps=steps)
    return forecast.predicted_mean, forecast.conf_int()

def fit_prophet(series, steps):
    prophet_df = series.reset_index().rename(columns={"Date": "ds", "CRYPTO": "y"})
    model = Prophet(daily_seasonality=True)
    model.fit(prophet_df)
    future = model.make_future_dataframe(periods=steps)
    forecast = model.predict(future)
    return forecast[['ds', 'yhat', 'yhat_lower', 'yhat_upper']].set_index('ds').iloc[-steps:]

def fit_lstm(series, steps, epochs=40, neurons=64):
    from tensorflow.keras.models import Sequential
    from tensorflow.keras.layers import LSTM, Dense
    scaler = MinMaxScaler()
    scaled = scaler.fit_transform(series.values.reshape(-1, 1))
    X, y = [], []
    look_back = 30
    for i in range(look_back, len(scaled)):
        X.append(scaled[i-look_back:i, 0])
        y.append(scaled[i, 0])
    X, y = np.array(X), np.array(y)
    X = np.reshape(X, (X.shape[0], X.shape[1], 1))

    model = Sequential()
    model.add(LSTM(neurons, input_shape=(X.shape[1], 1)))
    model.add(Dense(1))
    model.compile(optimizer='adam', loss='mean_squared_error')
    model.fit(X, y, epochs=epochs, batch_size=16, verbose=0)

    inputs = scaled[-look_back:]
    preds = []
    for _ in range(steps):
        input_reshaped = np.reshape(inputs, (1, look_back, 1))
        pred = model.predict(input_reshaped, verbose=0)
        preds.append(pred[0,0])
        inputs = np.append(inputs[1:], pred)

    preds = scaler.inverse_transform(np.array(preds).reshape(-1,1)).flatten()
    return preds

def predict_signal(series):
    returns = series.pct_change().dropna()
    X = returns.shift(1).dropna().values.reshape(-1,1)
    y = (returns[1:] > 0).astype(int).values
    model = LogisticRegression()
    model.fit(X, y)
    last_return = returns.iloc[-1]
    prediction = model.predict([[last_return]])
    return "BUY" if prediction[0]==1 else "SELL"

def get_sentiment_score(keyword="Bitcoin", n_posts=30):
    sia = SentimentIntensityAnalyzer()
    scores = []

    search_url = f"https://www.reddit.com/search/?q={keyword}"
    headers = {"User-Agent": "Mozilla/5.0"}
    response = requests.get(search_url, headers=headers)

    if response.status_code == 200:
        soup = BeautifulSoup(response.text, 'html.parser')
        posts = soup.find_all('h3', limit=n_posts)  # Titles usually inside <h3>
        for post in posts:
            text = post.get_text()
            sentiment = sia.polarity_scores(text)
            scores.append(sentiment['compound'])
    else:
        print("Failed to fetch Reddit posts. Status code:", response.status_code)

    return np.mean(scores) if scores else 0
def calculate_var(series, confidence_level=0.95):
    returns = series.pct_change().dropna()
    var = np.percentile(returns, (1-confidence_level)*100)
    return round(var * 100, 2)

def detect_anomalies(series):
    model = IsolationForest(contamination=0.01)
    preds = model.fit_predict(series.values.reshape(-1,1))
    anomalies = series[preds == -1]
    return anomalies

def classify_risk(volatility, var95, fragility, sentiment):
    score = (volatility*0.4 + abs(var95)*0.3 + (100-fragility)*0.2 + (1-sentiment)*0.1)
    if score > 80:
        return "🔥 Very High Risk"
    elif score > 60:
        return "⚡ High Risk"
    elif score > 40:
        return "⚠️ Moderate Risk"
    else:
        return "✅ Low Risk"

def backtest_model(series, model_choice, days_back=30):
    recent_series = series[-(days_back+60):-days_back]
    future_real = series[-days_back:]
    if model_choice == "ARIMA":
        pred, _ = fit_arima(recent_series, days_back)
    elif model_choice == "Prophet":
        pred = fit_prophet(recent_series, days_back)['yhat']
    else:
        pred = fit_lstm(recent_series, days_back)
    pred = pd.Series(pred, index=future_real.index)
    mse = ((future_real - pred)**2).mean()
    return round(mse, 4)

def generate_advanced_pptx(asset_name, summary_stats, signal, risk_score, shock_info, model_used, sentiment, var_95, var_99, risk_classification, backtest_mse):
    prs = Presentation()

    # ========== Colors (wrapped properly) ==========
    dark_blue = RGBColor(0, 51, 102)
    grey = RGBColor(102, 102, 102)
    black = RGBColor(0, 0, 0)
    light_blue_box = RGBColor(220, 230, 241)
    light_red = RGBColor(255, 102, 102)
    white = RGBColor(255, 255, 255)

    # ========== Cover Slide ==========
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)

    # Title
    title = slide.shapes.add_textbox(Inches(0.5), Inches(1), Inches(9), Inches(1))
    title_tf = title.text_frame
    p = title_tf.add_paragraph()
    p.text = f"{asset_name} - Stress Test Report"
    p.font.bold = True
    p.font.size = Pt(40)
    p.font.color.rgb = dark_blue

    # Subtitle
    subtitle = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(1))
    subtitle_tf = subtitle.text_frame
    p = subtitle_tf.add_paragraph()
    p.text = f"Generated on {datetime.datetime.today().strftime('%Y-%m-%d')}\nModel Used: {model_used}"
    p.font.size = Pt(20)
    p.font.color.rgb = grey

    # ========== Key Metrics Slide ==========
    slide = prs.slides.add_slide(slide_layout)
    metrics_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(5))
    tf = metrics_box.text_frame
    tf.word_wrap = True

    p = tf.add_paragraph()
    p.text = "📊 Key Metrics"
    p.font.bold = True
    p.font.size = Pt(30)
    p.space_after = Pt(20)

    for key, value in summary_stats.items():
        para = tf.add_paragraph()
        para.text = f"• {key}: {value}"
        para.font.size = Pt(22)
        para.font.color.rgb = black

    # ========== Additional Analysis Slide ==========
    slide = prs.slides.add_slide(slide_layout)
    analysis_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(5))
    tf = analysis_box.text_frame
    tf.word_wrap = True

    p = tf.add_paragraph()
    p.text = "🧠 Sentiment and Risk Analysis"
    p.font.bold = True
    p.font.size = Pt(30)
    p.space_after = Pt(20)

    fields = {
        "AI Trading Signal": signal,
        "Fragility Risk Score": f"{risk_score:.2f}/100",
        "Sentiment Score": f"{sentiment:.3f}",
        "VaR 95%": f"{var_95:.2f}%",
        "VaR 99%": f"{var_99:.2f}%",
        "Risk Classification": risk_classification,
        "Backtest MSE": backtest_mse
    }

    for key, value in fields.items():
        para = tf.add_paragraph()
        para.text = f"• {key}: {value}"
        para.font.size = Pt(22)
        para.font.color.rgb = black

    # ========== Shock Scenario Slide ==========
    slide = prs.slides.add_slide(slide_layout)

    shock_box = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(8.5), Inches(5))
    shock_box.fill.solid()
    shock_box.fill.fore_color.rgb = light_red
    shock_box.line.color.rgb = light_red

    tf = shock_box.text_frame
    tf.word_wrap = True

    p = tf.add_paragraph()
    p.text = "💥 Shock Scenario"
    p.font.bold = True
    p.font.size = Pt(30)
    p.font.color.rgb = white
    p.space_after = Pt(20)

    for key, value in shock_info.items():
        para = tf.add_paragraph()
        para.text = f"• {key}: {value}"
        para.font.size = Pt(20)
        para.font.color.rgb = white

    # Save
    pptx_io = io.BytesIO()
    prs.save(pptx_io)
    pptx_io.seek(0)
    return pptx_io


def simulate_monte_carlo(series, n_simulations=100, n_days=30):
    returns = series.pct_change().dropna()
    last_price = series.iloc[-1]
    mu = returns.mean()
    sigma = returns.std()

    simulations = np.zeros((n_days, n_simulations))
    for i in range(n_simulations):
        shock = np.random.normal(loc=mu, scale=sigma, size=n_days)
        price_series = [last_price]
        for s in shock:
            price_series.append(price_series[-1] * (1 + s))
        simulations[:, i] = price_series[1:]

    return simulations

def plot_monte_carlo(simulations, last_date):
    # Create future dates
    future_dates = pd.date_range(start=last_date + pd.Timedelta(days=1), periods=simulations.shape[0])

    # Compute statistics
    median = np.median(simulations, axis=1)
    p10 = np.percentile(simulations, 10, axis=1)
    p90 = np.percentile(simulations, 90, axis=1)

    fig = go.Figure()

    # Plot a sample of individual random paths (e.g., 10)
    idx = np.random.choice(simulations.shape[1], size=10, replace=False)
    for i in idx:
        fig.add_trace(go.Scatter(
            x=future_dates, y=simulations[:, i],
            mode='lines',
            line=dict(width=1, color='gray'),
            opacity=0.5,
            showlegend=False
        ))

    # Plot confidence interval area (optional, very light)
    fig.add_traces([
        go.Scatter(
            x=future_dates, y=p90,
            line=dict(width=0),
            mode='lines',
            showlegend=False,
            hoverinfo='skip'
        ),
        go.Scatter(
            x=future_dates, y=p10,
            line=dict(width=0),
            mode='lines',
            fill='tonexty',
            fillcolor='rgba(0,176,246,0.05)',  # Very light blue
            showlegend=False,
            hoverinfo='skip'
        )
    ])

    # Plot median path
    fig.add_trace(go.Scatter(
        x=future_dates, y=median,
        mode='lines',
        line=dict(color='blue', width=3),
        name='Median Forecast'
    ))

    # Plot 10% and 90% bounds
    fig.add_trace(go.Scatter(
        x=future_dates, y=p10,
        mode='lines',
        line=dict(color='red', dash='dash', width=2),
        name='10% Worst Case'
    ))

    fig.add_trace(go.Scatter(
        x=future_dates, y=p90,
        mode='lines',
        line=dict(color='green', dash='dash', width=2),
        name='90% Best Case'
    ))

    # Layout
    fig.update_layout(
        title="🎲 Monte Carlo Simulation - Selected Paths and Confidence",
        template=template,
        height=700,
        hovermode="x unified",
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1),
        xaxis_title="Date",
        yaxis_title="Simulated Price"
    )

    return fig

# ========== INTERFACE ========== #
# ========== INTERFACE (CONTINUED) ========== #

st.title("📈 Crypto Stress Test AI - Full Professional Intelligent Edition")
st.caption("Due Diligence AI | Advanced Forecasting, Risk Analysis & Reporting 🚀")
st.divider()

col1, col2, col3 = st.columns(3)
with col1:
    available_files = load_files(FOLDER_PATH)
    file_selected = st.selectbox("📁 Crypto Asset:", available_files)
with col2:
    model_choice = st.radio("🧠 Model:", ["ARIMA", "Prophet", "LSTM"], horizontal=True)
with col3:
    lookback_days = st.slider("📅 Lookback Days:", 60, 365, 180, step=30)

col4, col5, col6 = st.columns(3)
with col4:
    forecast_days = st.slider("🔮 Forecast Horizon (days):", 7, 90, 30)
with col5:
    shock_pct = st.slider("💥 Simulated Shock (%):", -90, 90, -30)
with col6:
    shock_day = st.slider("📆 Shock Day:", 10, 350, 70)

filepath = os.path.join(FOLDER_PATH, file_selected)
df = load_crypto_data(filepath, lookback_days)
if df.empty:
    st.error("❌ No valid data found.")
    st.stop()

# Stress Test
stress_df = df.copy()
stress_date = stress_df.index[min(shock_day, len(df)-1)]
stress_df.loc[stress_date:, 'CRYPTO'] *= (1 + shock_pct/100)

# Fragility
volatility = df['CRYPTO'].pct_change().std()
fragility_score = max(0, 100 - volatility * 1000)

# Sentiment Analysis
with st.spinner("Fetching Sentiment Analysis..."):
    sentiment_score = get_sentiment_score(file_selected.replace('.csv','').split('_')[0])

# VaR Calculation
var_95 = calculate_var(df['CRYPTO'], confidence_level=0.95)
var_99 = calculate_var(df['CRYPTO'], confidence_level=0.99)

# Forecast
with st.spinner("Training model..."):
    if model_choice == "ARIMA":
        forecast_orig, _ = fit_arima(df['CRYPTO'], forecast_days)
        forecast_stress, _ = fit_arima(stress_df['CRYPTO'], forecast_days)
    elif model_choice == "Prophet":
        forecast_orig = fit_prophet(df['CRYPTO'], forecast_days)
        forecast_stress = fit_prophet(stress_df['CRYPTO'], forecast_days)
    else:
        forecast_orig = fit_lstm(df['CRYPTO'], forecast_days)
        forecast_stress = fit_lstm(stress_df['CRYPTO'], forecast_days)

# Anomalies Detection
anomalies = detect_anomalies(df['CRYPTO'])

# Signal
signal = predict_signal(df['CRYPTO'])

# Risk Classification
risk_classification = classify_risk(volatility, var_95, fragility_score, sentiment_score)

# Forecast Backtest
mse_backtest = backtest_model(df['CRYPTO'], model_choice)

# Visualization
future_idx = pd.date_range(start=df.index[-1]+pd.Timedelta(days=1), periods=forecast_days, freq='D')
fig = go.Figure()
fig.add_trace(go.Scatter(x=df.index, y=df['CRYPTO'], name="Original", line=dict(color="#00BFFF")))
fig.add_trace(go.Scatter(x=stress_df.index, y=stress_df['CRYPTO'], name="Stress", line=dict(color="#FF4500")))
fig.add_trace(go.Scatter(x=anomalies.index, y=anomalies, mode='markers', name='Anomalies', marker=dict(size=8, color="red", symbol="x")))

if model_choice == "ARIMA":
    fig.add_trace(go.Scatter(x=future_idx, y=forecast_orig, name="Forecast Original", line=dict(dash='dot')))
    fig.add_trace(go.Scatter(x=future_idx, y=forecast_stress, name="Forecast Stress", line=dict(dash='dot')))
elif model_choice == "Prophet":
    fig.add_trace(go.Scatter(x=forecast_orig.index, y=forecast_orig['yhat'], name="Forecast Original", line=dict(dash='dot')))
    fig.add_trace(go.Scatter(x=forecast_stress.index, y=forecast_stress['yhat'], name="Forecast Stress", line=dict(dash='dot')))
else:
    fig.add_trace(go.Scatter(x=future_idx, y=forecast_orig, name="Forecast LSTM Original", line=dict(dash='dot')))
    fig.add_trace(go.Scatter(x=future_idx, y=forecast_stress, name="Forecast LSTM Stress", line=dict(dash='dot')))

fig.update_layout(template=template, hovermode="x unified", height=750, legend=dict(orientation="h", y=-0.2))
st.plotly_chart(fig, use_container_width=True)

# Prepare data for PPTX
summary_stats = {
    "Volatility": f"{volatility:.4f}",
    "Fragility Index": f"{fragility_score:.2f}/100",
    "Forecast Mean": f"{np.mean(forecast_orig) if not isinstance(forecast_orig, pd.DataFrame) else forecast_orig['yhat'].mean():.2f}",
    "Forecast Std Dev": f"{np.std(forecast_orig) if not isinstance(forecast_orig, pd.DataFrame) else forecast_orig['yhat'].std():.2f}"
}
shock_info = {
    "Shock Percentage": f"{shock_pct}%",
    "Shock Day": f"Day {shock_day} of time window"
}

# ====== Monte Carlo Simulation Section ======
with st.expander("🧪 Monte Carlo Stress Test (500 Simulations)"):
    with st.spinner("Running Monte Carlo Simulations..."):
        simulations = simulate_monte_carlo(df['CRYPTO'], n_simulations=500, n_days=forecast_days)
        fig_mc = plot_monte_carlo(simulations, df.index[-1])
        st.plotly_chart(fig_mc, use_container_width=True)

# ====== Intelligent Future Prediction Section ======
with st.expander("🤖 Intelligent AI Future Prediction"):
    st.subheader("🔮 Predict Future Prices (3 / 7 / 30 days)")
    prediction_days = st.selectbox("Select prediction horizon:", [3, 7, 30])

    with st.spinner("Training intelligent AI model..."):
        # Prepare Data
        scaler = MinMaxScaler()
        scaled_data = scaler.fit_transform(df['CRYPTO'].values.reshape(-1, 1))
        
        look_back = 30
        X, y = [], []
        for i in range(look_back, len(scaled_data)):
            X.append(scaled_data[i-look_back:i, 0])
            y.append(scaled_data[i, 0])
        X, y = np.array(X), np.array(y)
        X = X.reshape((X.shape[0], X.shape[1], 1))

        # Build LSTM Model
        from tensorflow.keras.models import Sequential
        from tensorflow.keras.layers import LSTM, Dense
        model = Sequential()
        model.add(LSTM(64, input_shape=(look_back, 1)))
        model.add(Dense(1))
        model.compile(optimizer='adam', loss='mse')
        model.fit(X, y, epochs=50, batch_size=16, verbose=0)

        # Predict future
        last_input = scaled_data[-look_back:]
        future_preds = []
        current_input = last_input
        for _ in range(prediction_days):
            pred = model.predict(current_input.reshape(1, look_back, 1), verbose=0)
            future_preds.append(pred[0,0])
            current_input = np.append(current_input[1:], pred)

        future_preds = scaler.inverse_transform(np.array(future_preds).reshape(-1, 1)).flatten()

        # Prepare dates
        last_date = df.index[-1]
        future_dates = pd.date_range(start=last_date + pd.Timedelta(days=1), periods=prediction_days)

        # Plot
        fig_pred = go.Figure()
        fig_pred.add_trace(go.Scatter(x=df.index, y=df['CRYPTO'], name="Historical", line=dict(color="lightblue")))
        fig_pred.add_trace(go.Scatter(x=future_dates, y=future_preds, name="Future Prediction", line=dict(color="orange", dash='dash')))

        fig_pred.update_layout(
            title="🤖 Intelligent AI Prediction",
            template=template,
            height=600,
            hovermode="x unified",
            legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1),
            xaxis_title="Date",
            yaxis_title="Crypto Price"
        )
        st.plotly_chart(fig_pred, use_container_width=True)

# ====== Intelligent AI Prediction Explainability Section ======
# ====== Intelligent AI Prediction Explainability Section ======
# ====== Intelligent AI Prediction Explainability Section ======
with st.expander("🧠 Explain AI Prediction (Feature Importance)"):
    st.info("Interpreting how past prices influenced the AI's prediction...")

    import shap
    import matplotlib.pyplot as plt
    import streamlit.components.v1 as components
    import tensorflow as tf

    # Step 1: Flatten input X
    X_flat = X.reshape((X.shape[0], X.shape[1]))  # (samples, 30)

    # Step 2: Create a Wrapper model
    class WrappedModel(tf.keras.Model):
        def __init__(self, original_model):
            super().__init__()
            self.original_model = original_model
        
        def call(self, inputs):
            inputs = tf.expand_dims(inputs, -1)  # (batch, 30) -> (batch, 30, 1)
            return self.original_model(inputs)

    wrapped_model = WrappedModel(model)

    # Step 3: SHAP KernelExplainer
    background = X_flat[:50]  # Background samples
    explainer = shap.KernelExplainer(wrapped_model.predict, background)

    # Step 4: Explain on test samples
    test_samples = X_flat[50:100]
    shap_values_list = explainer.shap_values(test_samples, nsamples=100)

    # ✅ Important fix:
    shap_values = shap_values_list if isinstance(shap_values_list, np.ndarray) else shap_values_list[0]

    # Step 5: Plot SHAP Summary
    st.subheader("🔍 SHAP Feature Importance (Past Lags)")
    shap_plot = plt.figure()
    shap.summary_plot(shap_values, test_samples, show=False, plot_type="bar", feature_names=[f"Day-{i}" for i in range(X_flat.shape[1])])
    st.pyplot(shap_plot)

    # Step 6: SHAP Force Plot for Single Sample
    st.subheader("⚡ SHAP Force Plot for a Random Prediction")

    random_idx = np.random.randint(0, len(test_samples))
    shap_single_value = shap_values[random_idx]

    shap_html = shap.force_plot(
        explainer.expected_value, 
        shap_single_value,
        feature_names=[f"Day-{i}" for i in range(X_flat.shape[1])],
        matplotlib=False
    )

    components.html(shap_html.html(), height=400)

pptx_file = generate_advanced_pptx(file_selected, summary_stats, signal, fragility_score, shock_info, model_choice, sentiment_score, var_95, var_99, risk_classification, mse_backtest)


# Export
st.download_button("💾 Download Full Intelligent PPTX Report", data=pptx_file, file_name="crypto_stress_intelligent_full_ai.pptx", mime="application/vnd.openxmlformats-officedocument.presentationml.presentation")

st.caption("Built with ❤️ | Full Advanced Crypto Risk Platform | Due Diligence AI 🚀")
