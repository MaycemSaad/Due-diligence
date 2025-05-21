# 🚀 Ultimate Crypto Stress Test & AI Report Generator

import streamlit as st
import pandas as pd
import numpy as np
import plotly.graph_objects as go
from sklearn.linear_model import LogisticRegression
from sklearn.preprocessing import MinMaxScaler
from tensorflow.keras.models import Sequential
from tensorflow.keras.layers import LSTM, Dense
from statsmodels.tsa.arima.model import ARIMA
from prophet import Prophet
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from io import BytesIO
import os

# ==================== CONFIG ====================
st.set_page_config(page_title="Crypto Stress Test PRO AI", layout="wide")

# --- Light/Dark mode ---
if "mode" not in st.session_state:
    st.session_state.mode = "Light"
mode_choice = st.toggle("🌗 Dark / Light Mode", value=(st.session_state.mode == "Dark"))
st.session_state.mode = "Dark" if mode_choice else "Light"
template = "plotly_dark" if st.session_state.mode == "Dark" else "plotly_white"

FOLDER_PATH = "C:/Users/informasud/Desktop/Data-Preparation/LOT3"

# ==================== FUNCTIONS ====================
@st.cache_data
def load_files(folder_path):
    return [f for f in os.listdir(folder_path) if f.endswith(".csv")]

@st.cache_data
def load_crypto_data(filepath, lookback_days):
    df = pd.read_csv(filepath)
    df.columns = [col.strip().replace('"', '') for col in df.columns]
    df['Date'] = pd.to_datetime(df['Date'], format='%m/%d/%Y', errors='coerce')
    df = df.dropna(subset=['Date'])
    df['Price'] = df['Price'].astype(str).str.replace(',', '').astype(float)
    df = df.sort_values('Date').set_index('Date')
    df = df[['Price']].rename(columns={"Price": "CRYPTO"})
    df = df.last(f'{lookback_days}D')
    return df

def fit_lstm(series, steps, epochs=40, neurons=64):
    scaler = MinMaxScaler()
    scaled = scaler.fit_transform(series.values.reshape(-1, 1))
    X, y = [], []
    look_back = 30
    for i in range(look_back, len(scaled)):
        X.append(scaled[i-look_back:i, 0])
        y.append(scaled[i, 0])
    X, y = np.array(X), np.array(y)
    X = np.reshape(X, (X.shape[0], X.shape[1], 1))

    model = Sequential()
    model.add(LSTM(neurons, input_shape=(X.shape[1], 1)))
    model.add(Dense(1))
    model.compile(loss='mean_squared_error', optimizer='adam')
    model.fit(X, y, epochs=epochs, batch_size=16, verbose=0)

    future_inputs = scaled[-look_back:]
    preds = []
    for _ in range(steps):
        input_reshaped = np.reshape(future_inputs, (1, look_back, 1))
        pred = model.predict(input_reshaped, verbose=0)
        preds.append(pred[0,0])
        future_inputs = np.append(future_inputs[1:], pred)

    preds_inverse = scaler.inverse_transform(np.array(preds).reshape(-1,1)).flatten()
    return preds_inverse

def predict_direction(series):
    df = pd.DataFrame(series)
    df['Target'] = (df['CRYPTO'].shift(-1) > df['CRYPTO']).astype(int)
    df = df.dropna()
    X = df[['CRYPTO']]
    y = df['Target']

    model = LogisticRegression()
    model.fit(X, y)
    latest = series.iloc[-1]
    prediction = model.predict([[latest]])[0]
    return prediction

def generate_pptx(summary_text, recommendation, chart_path):
    prs = Presentation()
    slide_layout = prs.slide_layouts[5]

    # Title Slide
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Crypto Stress Test AI Results"

    # Summary Slide
    slide = prs.slides.add_slide(slide_layout)
    tf = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(5)).text_frame
    p = tf.add_paragraph()
    p.text = summary_text
    p.font.size = Pt(18)

    # Signal Slide
    slide = prs.slides.add_slide(slide_layout)
    tf2 = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text_frame
    p2 = tf2.add_paragraph()
    p2.text = f"Signal Recommendation: {recommendation}"
    p2.font.size = Pt(24)
    p2.font.bold = True
    if recommendation == "BUY":
        p2.font.color.rgb = RGBColor(0, 176, 80)
    else:
        p2.font.color.rgb = RGBColor(255, 0, 0)

    # Chart Slide
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.add_picture(chart_path, Inches(1), Inches(1.5), Inches(7), Inches(4))

    pptx_io = BytesIO()
    prs.save(pptx_io)
    pptx_io.seek(0)
    return pptx_io

# ==================== APP ====================
st.title("📈 Crypto Stress Test - Deep AI Edition")
st.divider()

# --- Inputs ---
col1, col2, col3 = st.columns(3)
with col1:
    available_files = load_files(FOLDER_PATH)
    file_selected = st.selectbox("📁 Crypto à analyser:", available_files)
with col2:
    lookback_days = st.slider("⏳ Fenêtre temporelle:", 60, 365, 180)
with col3:
    forecast_days = st.slider("🔮 Jours de prédiction:", 7, 60, 14)

col4, col5 = st.columns(2)
with col4:
    shock_pct = st.slider("💥 Choc simulé:", -90, 90, -30)
with col5:
    model_choice = st.radio("🧠 Modèle:", ["LSTM"], horizontal=True)

# --- Load Data ---
filepath = os.path.join(FOLDER_PATH, file_selected)
df = load_crypto_data(filepath, lookback_days)
if df.empty:
    st.error("❌ Aucun fichier valide.")
    st.stop()

# --- Stress Test ---
stress_df = df.copy()
shock_day = min(30, len(df)-10)
stress_df.iloc[shock_day:, 0] *= (1 + shock_pct/100)

# --- Fragility ---
vol = df['CRYPTO'].rolling(window=30).std().mean()
fragility_score = max(0, 100 - vol * 1000)

# --- Forecasting ---
future_index = pd.date_range(df.index[-1] + pd.Timedelta(days=1), periods=forecast_days)
forecast_orig = fit_lstm(df['CRYPTO'], forecast_days)
forecast_stress = fit_lstm(stress_df['CRYPTO'], forecast_days)

# --- Direction Prediction ---
predicted_direction = predict_direction(df['CRYPTO'])
signal = "BUY" if predicted_direction == 1 else "SELL"

# --- Plotting ---
st.subheader("📊 Résultats Visuels")
fig, ax = plt.subplots(figsize=(10,6))
plt.plot(df.index, df['CRYPTO'], label='Original', color='blue')
plt.plot(stress_df.index, stress_df['CRYPTO'], label='Stressé', color='red')
plt.plot(future_index, forecast_orig, '--', label='Forecast Orig', color='green')
plt.plot(future_index, forecast_stress, '--', label='Forecast Stressé', color='orange')
plt.legend()
plt.grid()
chart_path = "temp_chart.png"
plt.savefig(chart_path)

st.plotly_chart(go.Figure(go.Scatter(x=df.index, y=df['CRYPTO'], name='Original')), use_container_width=True)

# --- Export PPTX ---
summary_text = f"Forecast: {forecast_days} jours | Fragilité: {fragility_score:.2f}/100"
pptx_file = generate_pptx(summary_text, signal, chart_path)

st.download_button(
    label="💾 Télécharger Présentation AI",
    data=pptx_file,
    file_name="forecast_ai_summary.pptx",
    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
)

st.caption("Made with ❤️ | Due Diligence AI | Deep Learning | LSTM | Logistic Regression")
