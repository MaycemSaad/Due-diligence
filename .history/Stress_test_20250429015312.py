# 🚀 Ultimate Crypto Stress Test AI - Full Professional Intelligent Version

import streamlit as st
import pandas as pd
import numpy as np
import plotly.graph_objects as go
from sklearn.preprocessing import MinMaxScaler
from sklearn.linear_model import LogisticRegression
from sklearn.ensemble import IsolationForest
from statsmodels.tsa.arima.model import ARIMA
from prophet import Prophet
from pptx import Presentation
from pptx.util import Inches, Pt
import snscrape.modules.reddit as reddit_scraper
from nltk.sentiment import SentimentIntensityAnalyzer
import io
import os
import datetime

# ========== CONFIGURATION ========== #
st.set_page_config(page_title="Crypto Stress Test AI - Ultimate Pro", layout="wide")

if "mode" not in st.session_state:
    st.session_state.mode = "Light"
mode_choice = st.toggle("🌗 Dark / Light Mode", value=(st.session_state.mode == "Dark"))
st.session_state.mode = "Dark" if mode_choice else "Light"
template = "plotly_dark" if st.session_state.mode == "Dark" else "plotly_white"

FOLDER_PATH = "C:/Users/informasud/Desktop/Data-Preparation/LOT3"

# ========== FUNCTIONS ========== #

@st.cache_data
def load_files(folder_path):
    return [f for f in os.listdir(folder_path) if f.endswith(".csv")]

@st.cache_data
def load_crypto_data(filepath, lookback_days):
    df = pd.read_csv(filepath)
    df.columns = [col.strip().replace('"', '') for col in df.columns]
    df['Date'] = pd.to_datetime(df['Date'], errors='coerce')
    df = df.dropna(subset=['Date'])
    df['Price'] = df['Price'].astype(str).str.replace(',', '').astype(float)
    df = df.sort_values('Date').set_index('Date')
    df = df[['Price']].rename(columns={"Price": "CRYPTO"})
    df = df.last(f'{lookback_days}D')
    return df

def fit_arima(series, steps):
    model = ARIMA(series, order=(2,1,2))
    fitted = model.fit()
    forecast = fitted.get_forecast(steps=steps)
    return forecast.predicted_mean, forecast.conf_int()

def fit_prophet(series, steps):
    prophet_df = series.reset_index().rename(columns={"Date": "ds", "CRYPTO": "y"})
    model = Prophet(daily_seasonality=True)
    model.fit(prophet_df)
    future = model.make_future_dataframe(periods=steps)
    forecast = model.predict(future)
    return forecast[['ds', 'yhat', 'yhat_lower', 'yhat_upper']].set_index('ds').iloc[-steps:]

def fit_lstm(series, steps, epochs=40, neurons=64):
    from tensorflow.keras.models import Sequential
    from tensorflow.keras.layers import LSTM, Dense
    scaler = MinMaxScaler()
    scaled = scaler.fit_transform(series.values.reshape(-1, 1))
    X, y = [], []
    look_back = 30
    for i in range(look_back, len(scaled)):
        X.append(scaled[i-look_back:i, 0])
        y.append(scaled[i, 0])
    X, y = np.array(X), np.array(y)
    X = np.reshape(X, (X.shape[0], X.shape[1], 1))

    model = Sequential()
    model.add(LSTM(neurons, input_shape=(X.shape[1], 1)))
    model.add(Dense(1))
    model.compile(optimizer='adam', loss='mean_squared_error')
    model.fit(X, y, epochs=epochs, batch_size=16, verbose=0)

    inputs = scaled[-look_back:]
    preds = []
    for _ in range(steps):
        input_reshaped = np.reshape(inputs, (1, look_back, 1))
        pred = model.predict(input_reshaped, verbose=0)
        preds.append(pred[0,0])
        inputs = np.append(inputs[1:], pred)

    preds = scaler.inverse_transform(np.array(preds).reshape(-1,1)).flatten()
    return preds

def predict_signal(series):
    returns = series.pct_change().dropna()
    X = returns.shift(1).dropna().values.reshape(-1,1)
    y = (returns[1:] > 0).astype(int).values
    model = LogisticRegression()
    model.fit(X, y)
    last_return = returns.iloc[-1]
    prediction = model.predict([[last_return]])
    return "BUY" if prediction[0]==1 else "SELL"

def get_sentiment_score(keyword="Bitcoin", n_posts=50):
    sia = SentimentIntensityAnalyzer()
    scores = []
    for post in reddit_scraper.RedditSearchScraper(f'{keyword}').get_items():
        if len(scores) >= n_posts:
            break
        sentiment = sia.polarity_scores(post.content)
        scores.append(sentiment['compound'])
    return np.mean(scores) if scores else 0

def calculate_var(series, confidence_level=0.95):
    returns = series.pct_change().dropna()
    var = np.percentile(returns, (1-confidence_level)*100)
    return round(var * 100, 2)

def detect_anomalies(series):
    model = IsolationForest(contamination=0.01)
    preds = model.fit_predict(series.values.reshape(-1,1))
    anomalies = series[preds == -1]
    return anomalies

def classify_risk(volatility, var95, fragility, sentiment):
    score = (volatility*0.4 + abs(var95)*0.3 + (100-fragility)*0.2 + (1-sentiment)*0.1)
    if score > 80:
        return "🔥 Very High Risk"
    elif score > 60:
        return "⚡ High Risk"
    elif score > 40:
        return "⚠️ Moderate Risk"
    else:
        return "✅ Low Risk"

def backtest_model(series, model_choice, days_back=30):
    recent_series = series[-(days_back+60):-days_back]
    future_real = series[-days_back:]
    if model_choice == "ARIMA":
        pred, _ = fit_arima(recent_series, days_back)
    elif model_choice == "Prophet":
        pred = fit_prophet(recent_series, days_back)['yhat']
    else:
        pred = fit_lstm(recent_series, days_back)
    pred = pd.Series(pred, index=future_real.index)
    mse = ((future_real - pred)**2).mean()
    return round(mse, 4)

def generate_advanced_pptx(asset_name, summary_stats, signal, risk_score, shock_info, model_used, sentiment, var_95, var_99, risk_classification, backtest_mse):
    prs = Presentation()

    # ========== Cover Slide ========== 
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = f"{asset_name} - Stress Test Report"
    slide.placeholders[1].text = f"Generated on {datetime.datetime.today().strftime('%Y-%m-%d')}\nModel Used: {model_used}"

    # ========== Statistics Slide ========== 
    slide_layout = prs.slide_layouts[5]
    slide = prs.slides.add_slide(slide_layout)
    tf = slide.shapes.add_textbox(Inches(0.7), Inches(1), Inches(8), Inches(5)).text_frame
    tf.word_wrap = True
    tf.text = f"🔹 Key Metrics:\n\n"
    for key, value in summary_stats.items():
        p = tf.add_paragraph()
        p.text = f"• {key}: {value}"
        p.font.size = Pt(18)

    # ========== Extra Metrics Slide ========== 
    slide = prs.slides.add_slide(slide_layout)
    tf = slide.shapes.add_textbox(Inches(0.7), Inches(1), Inches(8), Inches(5)).text_frame
    tf.word_wrap = True
    tf.text = f"🧠 Additional Analysis:\n\n"
    for key, value in {
        "Sentiment Score": f"{sentiment:.3f}",
        "VaR 95%": f"{var_95:.2f}%",
        "VaR 99%": f"{var_99:.2f}%",
        "Risk Classification": risk_classification,
        "Backtest MSE": backtest_mse
    }.items():
        p = tf.add_paragraph()
        p.text = f"• {key}: {value}"
        p.font.size = Pt(18)

    # ========== Shock Scenario Slide ========== 
    slide = prs.slides.add_slide(slide_layout)
    tf = slide.shapes.add_textbox(Inches(0.7), Inches(1), Inches(8), Inches(5)).text_frame
    tf.text = f"💥 Shock Scenario:\n\n"
    for key, value in shock_info.items():
        p = tf.add_paragraph()
        p.text = f"• {key}: {value}"
        p.font.size = Pt(18)

    # Save
    pptx_io = io.BytesIO()
    prs.save(pptx_io)
    pptx_io.seek(0)
    return pptx_io

# ========== INTERFACE ========== #
