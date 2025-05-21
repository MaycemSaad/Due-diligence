import streamlit as st
import pandas as pd
import numpy as np
import plotly.graph_objects as go
import matplotlib.pyplot as plt
import yfinance as yf
import os
import time
from datetime import datetime, timedelta
import pytz
import warnings
from typing import Dict, Tuple, Optional
import glob
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import logging
from logging.handlers import RotatingFileHandler
import json
from pathlib import Path

warnings.filterwarnings('ignore')
plt.style.use('seaborn-v0_8')
st.set_page_config(layout="wide", page_title="Crypto Intelligence Platform")

CRYPTO_MAPPING = {
    "BTC-USD": "Bitcoin",
    "ETH-USD": "Ethereum",
    "BNB-USD": "Binance Coin",
    "SOL-USD": "Solana",
    "XRP-USD": "Ripple",
    "ADA-USD": "Cardano",
    "DOGE-USD": "Dogecoin",
    "DOT-USD": "Polkadot",
    "AVAX-USD": "Avalanche",
    "MATIC-USD": "Polygon"
}

class CryptoDashboard:
    def __init__(self):
        self.current_ticker = None
        self.current_data = None

    def generate_pptx_report(self, ticker: str) -> str:
        try:
            prs = Presentation()
            now = datetime.now().strftime('%Y-%m-%d %H:%M')
            report_dir = "reports"
            os.makedirs(report_dir, exist_ok=True)
            output_path = os.path.join(report_dir, f"{CRYPTO_MAPPING[ticker]}_Analysis_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pptx")

            def add_slide_with_title(prs, title_text, subtitle_text=None, layout_idx=1):
                slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
                slide.shapes.title.text = title_text
                if subtitle_text and len(slide.placeholders) > 1:
                    slide.placeholders[1].text = subtitle_text
                return slide

            title_slide = add_slide_with_title(
                prs,
                f"{CRYPTO_MAPPING[ticker]} Analysis Report",
                f"Generated on {now}\nCrypto Intelligence Platform",
                layout_idx=0
            )

            try:
                logo_path = "logo.png"
                if os.path.exists(logo_path):
                    left = Inches(7.5)
                    top = Inches(0.5)
                    height = Inches(1)
                    title_slide.shapes.add_picture(logo_path, left, top, height=height)
            except Exception:
                pass

            summary_slide = add_slide_with_title(prs, "Executive Summary")
            content = summary_slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(4))
            tf = content.text_frame
            tf.word_wrap = True

            p = tf.add_paragraph()
            p.text = f"Comprehensive analysis of {CRYPTO_MAPPING[ticker]} market performance, technical indicators, and sentiment analysis."
            p.font.size = Pt(14)

            p = tf.add_paragraph()
            p.text = "Key Findings:"
            p.font.size = Pt(14)
            p.font.bold = True

            bullet_points = [
                "Current market trends and price action",
                "Technical indicators and trading signals",
                "Sentiment analysis from news and social media",
                "Risk metrics and volatility analysis"
            ]
            for point in bullet_points:
                p = tf.add_paragraph()
                p.text = point
                p.font.size = Pt(12)
                p.level = 1

            price_slide = add_slide_with_title(prs, "Price Performance")
            fig = go.Figure()
            fig.add_trace(go.Scatter(x=self.current_data.index, y=self.current_data['Close'], name="Price", line=dict(color='#1f77b4', width=2)))

            if 'MA_20' in self.current_data.columns:
                fig.add_trace(go.Scatter(x=self.current_data.index, y=self.current_data['MA_20'], name="20-Day MA", line=dict(color='orange', dash='dot', width=1.5)))

            if 'MA_50' in self.current_data.columns:
                fig.add_trace(go.Scatter(x=self.current_data.index, y=self.current_data['MA_50'], name="50-Day MA", line=dict(color='green', dash='dot', width=1.5)))

            fig.update_layout(
                title=f"{CRYPTO_MAPPING[ticker]} Price History",
                xaxis_title="Date",
                yaxis_title="Price (USD)",
                template="plotly_white",
                height=500,
                margin=dict(l=20, r=20, t=60, b=20)
            )

            chart_path = f"temp_{ticker}_price_chart.png"
            fig.write_image(chart_path, width=1000, height=500, scale=2)
            price_slide.shapes.add_picture(chart_path, Inches(0.5), Inches(1.5), width=Inches(9), height=Inches(4.5))

            tech_slide = add_slide_with_title(prs, "Technical Indicators")
            left_pos = [0.5, 5.0]
            top_pos = [1.5, 3.5]

            if 'RSI' in self.current_data.columns:
                fig_rsi = go.Figure()
                fig_rsi.add_trace(go.Scatter(x=self.current_data.index, y=self.current_data['RSI'], name="RSI", line=dict(color='purple', width=2)))
                fig_rsi.add_hline(y=70, line_dash="dash", line_color="red")
                fig_rsi.add_hline(y=30, line_dash="dash", line_color="green")
                fig_rsi.update_layout(title="Relative Strength Index (RSI)", height=300, margin=dict(l=20, r=20, t=40, b=20))
                rsi_path = f"temp_{ticker}_rsi.png"
                fig_rsi.write_image(rsi_path, width=450, height=300, scale=2)
                tech_slide.shapes.add_picture(rsi_path, Inches(left_pos[0]), Inches(top_pos[0]), width=Inches(4), height=Inches(2.5))

            if all(col in self.current_data.columns for col in ['MACD', 'MACD_signal']):
                fig_macd = go.Figure()
                fig_macd.add_trace(go.Scatter(x=self.current_data.index, y=self.current_data['MACD'], name="MACD", line=dict(color='blue', width=1.5)))
                fig_macd.add_trace(go.Scatter(x=self.current_data.index, y=self.current_data['MACD_signal'], name="Signal", line=dict(color='orange', width=1.5)))
                fig_macd.update_layout(title="MACD", height=300, margin=dict(l=20, r=20, t=40, b=20))
                macd_path = f"temp_{ticker}_macd.png"
                fig_macd.write_image(macd_path, width=450, height=300, scale=2)
                tech_slide.shapes.add_picture(macd_path, Inches(left_pos[1]), Inches(top_pos[0]), width=Inches(4), height=Inches(2.5))

            if all(col in self.current_data.columns for col in ['BB_upper', 'BB_lower']):
                fig_bb = go.Figure()
                fig_bb.add_trace(go.Scatter(x=self.current_data.index, y=self.current_data['BB_upper'], name="Upper Band", line=dict(color='gray', width=1)))
                fig_bb.add_trace(go.Scatter(x=self.current_data.index, y=self.current_data['BB_lower'], name="Lower Band", line=dict(color='gray', width=1), fill='tonexty'))
                fig_bb.add_trace(go.Scatter(x=self.current_data.index, y=self.current_data['Close'], name="Price", line=dict(color='blue', width=1.5)))
                fig_bb.update_layout(title="Bollinger Bands", height=300, margin=dict(l=20, r=20, t=40, b=20))
                bb_path = f"temp_{ticker}_bb.png"
                fig_bb.write_image(bb_path, width=450, height=300, scale=2)
                tech_slide.shapes.add_picture(bb_path, Inches(left_pos[0]), Inches(top_pos[1]), width=Inches(4), height=Inches(2.5))

            if 'Volume' in self.current_data.columns:
                fig_vol = go.Figure()
                fig_vol.add_trace(go.Bar(x=self.current_data.index, y=self.current_data['Volume'], name="Volume", marker_color='#1f77b4'))
                fig_vol.update_layout(title="Trading Volume", height=300, margin=dict(l=20, r=20, t=40, b=20))
                vol_path = f"temp_{ticker}_vol.png"
                fig_vol.write_image(vol_path, width=450, height=300, scale=2)
                tech_slide.shapes.add_picture(vol_path, Inches(left_pos[1]), Inches(top_pos[1]), width=Inches(4), height=Inches(2.5))

            # Tu peux ajouter ici les parties suivantes de ton code (Slides : Signals, Status, Recommendations, Disclaimer...)

            prs.save(output_path)
            for f in glob.glob(f"temp_{ticker}_*.png"):
                try:
                    os.remove(f)
                except Exception:
                    pass
            return output_path

        except Exception as e:
            st.error(f"Error generating report: {str(e)}")
            raise

def main():
    dashboard = CryptoDashboard()
    dashboard.current_ticker = "BTC-USD"
    dashboard.current_data = yf.download(dashboard.current_ticker, period="1y")

    if st.button("Generate PowerPoint Report"):
        with st.spinner("Creating report..."):
            try:
                report_path = dashboard.generate_pptx_report(dashboard.current_ticker)
                st.success("Report generated successfully!")
                with open(report_path, "rb") as f:
                    st.download_button(
                        label="Download Report",
                        data=f,
                        file_name=os.path.basename(report_path),
                        mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
                    )
            except Exception as e:
                st.error(f"Failed to generate report: {str(e)}")

if __name__ == "__main__":
    main()
